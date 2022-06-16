import wps_utils as wu
import openpyxl
import asposecells
import jpype
from asposecells.api import *
import pickle
import os
import random as r
import win32com.client as win32

filepath = "C:/Users/gao/Desktop/SSP-四模-S5.xlsx"
name = "s5-2022"
wu.code_dst_gen(1, 6, name)
if os.path.exists("F:/py_xw/"+name+"_1.pkl"):
    dst = pickle.load(open(name+"_1.pkl", "rb"))
    dst = list(dst)
else:
    dst = pickle.load(open(name + ".pkl", "rb"))
    dst = list(dst)

x = r.randint(pow(10, 5),pow(10, 6))
while x in dst:
    x = r.randint(pow(10, 5),pow(10, 6))
dst.append(x)

wb1 = openpyxl.load_workbook(filepath)
ws_num = len(wb1.sheetnames)
print(ws_num)
wb = Workbook(filepath)

for ws_id in range(1, ws_num):
    row = 1
    ws = wb1.worksheets[ws_id]
    while (ws['B' + str(row)].value) is not None:
        row = row + 1
    wb = wu.add_watermark_wordart(wb, x, ws_id, row)
    # print(ws_id)
    # make_watermark(num, "tst", x_step=250, y_step=200, size=70)
    # add_pic(wb, "tst.png")
    # wb.save(filepath)
wb.save(filepath)
del wb
print(filepath)

# delete the last sheet
excel = win32.gencache.EnsureDispatch('Excel.Application')
excel.Visible = False
excel.DisplayAlerts = False
wb2 = excel.Workbooks.Open(filepath)
i = wb2.Worksheets.Count
wb2.Worksheets(i).Delete()
wb2.Save()
excel.DisplayAlerts = True
excel.Application.Quit()

pickle.dump(list, open(name + "_1.pkl", 'wb'))